# Rebuilds tools/test/data/sheets.json from Excel downloads of the season sheets (File → Download → .xlsx in Google Sheets).
#   python3 tools/test/xlsx2sheets.py <2025.xlsx> <2026.xlsx>
# Writes what the Sheets API would send the site: the same ranges the site asks for (Winnings A1:Q400, Rosters A1:AG10,
# QBs A1:A50), cells as displayed text ("$5", "+15", "0.5"), trailing empty cells and rows left off.
# Needs: pip install openpyxl
import sys, os, json, datetime
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
RANGES = {'Winnings': (400, 17), 'Rosters': (10, 33), 'QBs': (50, 1)}  # rows, columns

def show(c):
    v = c.value
    if v is None: return ''
    if isinstance(v, bool): return 'TRUE' if v else 'FALSE'
    if isinstance(v, (datetime.datetime, datetime.date)): return v.strftime('%-m/%-d/%Y %H:%M:%S') if isinstance(v, datetime.datetime) else v.strftime('%-m/%-d/%Y')
    if isinstance(v, (int, float)):
        n = int(v) if float(v).is_integer() else round(v, 4)
        if '$' in (c.number_format or ''): return ('-' if n < 0 else '') + '$' + format(abs(n), ',')
        return str(n)
    return str(v)

def tab(wb, name):
    rows_max, cols = RANGES[name]
    out = []
    for r in wb[name].iter_rows(min_row=1, max_row=rows_max, max_col=cols):
        vals = [show(c) for c in r]
        while vals and vals[-1] == '': vals.pop()
        out.append(vals)
    while out and not out[-1]: out.pop()
    return out

def main(f25, f26):
    w25 = openpyxl.load_workbook(f25, data_only=True)
    w26 = openpyxl.load_workbook(f26, data_only=True)
    D = {'Rosters': tab(w26, 'Rosters'), 'QBs': tab(w26, 'QBs'), '2025': tab(w25, 'Winnings'), '2026': tab(w26, 'Winnings')}
    out = os.path.join(HERE, 'data', 'sheets.json')
    json.dump(D, open(out, 'w'))
    print('wrote', out, {k: len(v) for k, v in D.items()})

if __name__ == '__main__':
    if len(sys.argv) != 3: sys.exit('usage: python3 tools/test/xlsx2sheets.py <2025.xlsx> <2026.xlsx>')
    main(sys.argv[1], sys.argv[2])
